"""Build a pixel-perfect PPTX from the 2x slide renders (full-bleed images)."""

from pptx import Presentation
from pptx.util import Inches

RENDER = "/home/z/my-project/download/render"
OUT = "/home/z/my-project/download/CyberShield-Presentation.pptx"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]  # fully blank layout

for i in range(1, 14):
    slide = prs.slides.add_slide(blank)
    slide.shapes.add_picture(
        f"{RENDER}/slide_{i:02d}.png", 0, 0,
        width=prs.slide_width, height=prs.slide_height,
    )

core = prs.core_properties
core.title = "CyberShield — Live Threat Intelligence"
core.author = "Umair Jamal"
core.subject = "Real-time cyber threat intelligence dashboard — project deck"

prs.save(OUT)
print("saved", OUT)
